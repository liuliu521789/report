# -*- coding: utf-8 -*-
"""Fix #REF! formulas in Sheet2 of 中来标签 - PR02.xlsx"""
from __future__ import annotations

import glob
import os
import sys


def find_src() -> str:
    desk = os.path.join(os.environ["USERPROFILE"], "Desktop")
    files = [
        p
        for p in glob.glob(os.path.join(desk, "*PR02*.xlsx"))
        if not os.path.basename(p).startswith("~$")
        and "fixed" not in os.path.basename(p).lower()
    ]
    if not files:
        raise FileNotFoundError("Desktop 上未找到 PR02 xlsx")
    return files[0]


def qr_formula(data_r: int, h_r: int) -> str:
    """二维码：用数据行 data_r 的字段 + H{h_r} 序列号。"""
    return (
        f'="M"&I{data_r}&"*B"&H{h_r}&"*G"&J{data_r}&"*C"&K{data_r}'
        f'&"*H"&L{data_r}&"*P"&M{data_r}&"*S"&N{data_r}'
        f'&"*T"&O{data_r}&"*D"&P{data_r}'
    )


def apply_fixes(set_formula):
    """
    规律（与未损坏行一致）：
      H{r} = CONCATENATE(J{r-1},Q{r-1},R{r-1})   # r>=3；H2 特例用同行
      G{r} 引用 I/J/K/... 的 {r-1} 与 H{r}
      标签：品名 S{r-1}、序列号 H{r}、数量 N{r-1}、批号 M{r-1}、编码 J{r-1}
    第 14 行断裂后，原 H15:G21 / 对应标签整体偏了一行，一并纠偏。
    """
    # --- row 3 ---
    set_formula("H3", "=CONCATENATE(J2,Q2,R2)")
    set_formula("G3", qr_formula(2, 3))
    set_formula("D9", "=S2")
    set_formula("D11", "=N2")
    set_formula("D12", "=M2")
    set_formula("D13", "=J2")

    # --- row 14 ---
    set_formula("H14", "=CONCATENATE(J13,Q13,R13)")
    set_formula("G14", qr_formula(13, 14))
    set_formula("D83", "=S13")
    set_formula("D85", "=N13")
    set_formula("D86", "=M13")
    set_formula("D87", "=J13")

    # --- realign 15..21 (were off-by-one) ---
    for r in range(15, 22):
        data_r = r - 1
        set_formula(f"H{r}", f"=CONCATENATE(J{data_r},Q{data_r},R{data_r})")
        set_formula(f"G{r}", qr_formula(data_r, r))

    # label serial cell row -> H row
    label_serial_rows = {
        90: 15,
        96: 16,
        102: 17,
        108: 18,
        114: 19,
        120: 20,
        126: 21,
    }
    for d_serial_row, h_r in label_serial_rows.items():
        data_r = h_r - 1
        set_formula(f"D{d_serial_row - 1}", f"=S{data_r}")
        set_formula(f"D{d_serial_row}", f"=H{h_r}")
        set_formula(f"D{d_serial_row + 1}", f"=N{data_r}")
        set_formula(f"D{d_serial_row + 2}", f"=M{data_r}")
        set_formula(f"D{d_serial_row + 3}", f"=J{data_r}")


def fix_via_com(src: str, out: str) -> bool:
    try:
        import win32com.client  # type: ignore
    except ImportError:
        print("pywin32 不可用，改用 openpyxl")
        return False

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    wb = None
    try:
        wb = excel.Workbooks.Open(src)
        ws = None
        for sh in wb.Worksheets:
            name = str(sh.Name)
            if "批次序列" in name or name.strip().startswith("2"):
                ws = sh
                break
        if ws is None:
            ws = wb.Worksheets(2)
        print("COM sheet:", ws.Name)

        def set_formula(addr: str, formula: str):
            cell = ws.Range(addr)
            cell.ClearContents()
            cell.NumberFormat = "General"
            cell.Formula = formula

        apply_fixes(set_formula)
        excel.CalculateFullRebuild()

        if os.path.exists(out):
            os.remove(out)
        wb.SaveAs(out)
        print("saved copy:", out)

        try:
            wb.SaveAs(src)
            print("also overwritten original")
        except Exception as e:
            print("original not overwritten (可能正被 Excel 打开):", e)
        return True
    finally:
        if wb is not None:
            wb.Close(SaveChanges=False)
        excel.Quit()


def fix_via_openpyxl(src: str, out: str) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(src)
    ws = wb[wb.sheetnames[1]]
    print("openpyxl sheet:", ws.title)

    def set_formula(addr: str, formula: str):
        ws[addr] = formula

    apply_fixes(set_formula)
    wb.save(out)
    print("saved copy:", out)
    try:
        wb.save(src)
        print("also overwritten original")
    except Exception as e:
        print("original not overwritten:", e)


def verify(path: str) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb[wb.sheetnames[1]]
    bad = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            v = cell.value
            if isinstance(v, str) and "#REF" in v:
                bad.append(f"{cell.coordinate}: {v}")
    if bad:
        print("仍有 #REF! :")
        for line in bad:
            print(" ", line)
        sys.exit(2)
    print("校验通过：Sheet2 公式中已无 #REF!")
    print("抽样:")
    for addr in ("G3", "H3", "D9", "G14", "H14", "D83", "G15", "H15", "D89", "G21", "H21"):
        print(f"  {addr} = {ws[addr].value}")


def main():
    src = find_src()
    desk = os.path.dirname(src)
    out = os.path.join(desk, "中来标签 - PR02-fixed.xlsx")
    print("src:", src)
    ok = False
    try:
        ok = fix_via_com(src, out)
    except Exception as e:
        print("COM 失败:", type(e).__name__, e)
    if not ok:
        fix_via_openpyxl(src, out)
    verify(out)


if __name__ == "__main__":
    main()
